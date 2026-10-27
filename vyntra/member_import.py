"""
Import members from an Excel (.xlsx) or CSV file, so a gym can move its register into Manaze in one go.

Each row becomes a member on one of the gym's plans, optionally with what they've paid. The app first sends the
file with `dry_run` to show which rows are fine and what's wrong with the rest, then imports the good rows.
"""

import csv
import io
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from .models import GymMember, Member, MembershipPlan
from .serializers import MemberSerializer
from .services import ActivityService, MembershipService, NotificationService, PaymentService

MAX_ROWS = 2000

# Column heading (lower-case, without * or extra spaces) → field. Several spellings are accepted.
COLUMNS = {
    "member id": "member_code", "member code": "member_code", "id": "member_code",
    "first name": "first_name", "name": "first_name", "full name": "first_name",
    "last name": "last_name", "surname": "last_name",
    "mobile": "phone", "mobile no": "phone", "mobile number": "phone", "phone": "phone",
    "email": "email",
    "gender": "gender",
    "date of birth": "date_of_birth", "dob": "date_of_birth",
    "plan": "plan", "plan name": "plan", "membership plan": "plan",
    "start date": "start_date", "joining date": "start_date",
    "end date": "end_date", "expiry date": "end_date",
    "discount": "discount",
    "amount paid": "amount_paid", "paid": "amount_paid",
    "payment method": "payment_method",
    "emergency contact name": "emergency_contact_name",
    "emergency contact mobile": "emergency_contact_phone",
    "address": "emergency_contact_address",
}

# Column names shown in error messages, for fields the serializer reports.
LABELS = {
    "member_code": "Member ID", "first_name": "First Name", "last_name": "Last Name", "phone": "Mobile",
    "email": "Email", "gender": "Gender", "date_of_birth": "Date of Birth", "membership_plan_id": "Plan",
    "start_date": "Start Date", "discount": "Discount", "payment": "Amount Paid",
    "emergency_contact_name": "Emergency Contact Name", "emergency_contact_phone": "Emergency Contact Mobile",
    "emergency_contact_address": "Address",
}

TEMPLATE_HEADINGS = [
    "Member ID", "First Name*", "Last Name", "Mobile", "Email", "Gender", "Date of Birth", "Plan*", "Start Date*",
    "End Date", "Discount", "Amount Paid", "Payment Method", "Emergency Contact Name", "Emergency Contact Mobile",
    "Address",
]
TEMPLATE_EXAMPLE = [
    "101", "Rahul", "Sharma", "9876543210", "rahul@example.com", "Male", "15-08-1995", "Monthly", "01-10-2026",
    "", "0", "1000", "UPI", "Priya Sharma", "9876500000", "Koramangala, Bengaluru",
]

GENDERS = {"m": "Male", "male": "Male", "f": "Female", "female": "Female", "other": "Other", "o": "Other"}
METHODS = {"cash": "CASH", "upi": "UPI", "card": "CARD", "bank": "BANK_TRANSFER", "bank transfer": "BANK_TRANSFER",
           "neft": "BANK_TRANSFER", "imps": "BANK_TRANSFER", "gpay": "UPI", "phonepe": "UPI", "paytm": "UPI"}
DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y", "%d %b %Y", "%d %B %Y")


class ImportFileError(Exception):
    """The file itself can't be read (wrong type, no headings, too many rows)."""


def template_csv():
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(TEMPLATE_HEADINGS)
    writer.writerow(TEMPLATE_EXAMPLE)
    return out.getvalue()


def _heading(value):
    return re.sub(r"\s+", " ", str(value or "").replace("*", "").replace(".", "").strip().lower())


def read_rows(upload):
    """[(row number, {field: raw value})] from an .xlsx or .csv upload."""
    name = (upload.name or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise ImportFileError("Couldn't open this Excel file. Save it as .xlsx and try again.")
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text)))
    else:
        raise ImportFileError("Upload an Excel (.xlsx) or CSV file.")

    while table and not any(str(cell or "").strip() for cell in table[0]):
        table.pop(0)
    if not table:
        raise ImportFileError("The file is empty.")
    fields = [COLUMNS.get(_heading(cell)) for cell in table[0]]
    if "first_name" not in fields or "plan" not in fields:
        raise ImportFileError("The first row must have the column headings, including First Name and Plan. "
                              "Download the template to see the format.")

    rows = []
    for number, cells in enumerate(table[1:], start=2):
        values = {field: cells[i] for i, field in enumerate(fields) if field and i < len(cells)}
        if any(str(v).strip() for v in values.values() if v is not None):
            rows.append((number, values))
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Import up to {MAX_ROWS} members at a time; this file has {len(rows)}.")
    return rows


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 9876543210 and 101 as floats
    return str(value).strip()


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"“{text}” isn't a date. Use DD-MM-YYYY, e.g. 01-10-2026.")


def _amount(value):
    text = _text(value).replace(",", "").replace("₹", "").strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(f"“{_text(value)}” isn't an amount.")


def _flatten(errors, prefix=""):
    """DRF error dict → ["Mobile: Enter a valid 10-digit mobile number."]"""
    messages = []
    for field, value in errors.items():
        label = LABELS.get(field, field.replace("_", " ").capitalize()) if field != "non_field_errors" else ""
        if isinstance(value, dict):
            messages += _flatten(value, label)
        else:
            for message in value if isinstance(value, list) else [value]:
                name = label or prefix
                messages.append(f"{name}: {message}" if name else str(message))
    return messages


def _next_codes(business_account_id):
    """Member IDs for rows without one: numbers after the highest numeric ID already used."""
    codes = Member.objects.filter(business_account_id=business_account_id).values_list("member_code", flat=True)
    highest = max((int(code) for code in codes if code and code.isdigit()), default=100)
    while True:
        highest += 1
        yield str(highest)


def check_rows(business_account_id, rows):
    """Validates every row. Returns [{row, ok, errors, name, member_code, plan, data, end_date}]."""
    plans = {plan.name.strip().lower(): plan for plan in MembershipPlan.objects.filter(
        business_account_id=business_account_id, is_active=True)}
    taken_codes = set(Member.objects.filter(business_account_id=business_account_id)
                      .values_list("member_code", flat=True))
    codes_in_file = {_text(v.get("member_code")) for _, v in rows if _text(v.get("member_code"))}
    new_codes = _next_codes(business_account_id)
    seen_codes = set()
    results = []

    for number, raw in rows:
        errors = []
        code = _text(raw.get("member_code"))
        if not code:
            code = next(c for c in new_codes if c not in taken_codes and c not in codes_in_file)
        if code in seen_codes:
            errors.append(f"Member ID: {code} appears more than once in the file.")
        seen_codes.add(code)

        first_name, last_name = _text(raw.get("first_name")), _text(raw.get("last_name"))
        if first_name and not last_name and " " in first_name and "last_name" not in raw:
            first_name, last_name = first_name.split(" ", 1)  # a single "Name" column

        plan_name = _text(raw.get("plan"))
        plan = plans.get(plan_name.lower())
        if not plan_name:
            errors.append("Plan: Enter the plan name.")
        elif plan is None:
            errors.append(f"Plan: No active plan called “{plan_name}”. Add it under Membership Plans first, "
                          f"or use one of: {', '.join(p.name for p in plans.values()) or 'none yet'}.")

        data = {
            "member_code": code, "first_name": first_name, "last_name": last_name or None,
            "phone": _text(raw.get("phone")) or None, "email": _text(raw.get("email")) or None,
            "gender": GENDERS.get(_text(raw.get("gender")).lower()) or (_text(raw.get("gender")) or None),
            "emergency_contact_name": _text(raw.get("emergency_contact_name")) or None,
            "emergency_contact_phone": _text(raw.get("emergency_contact_phone")) or None,
            "emergency_contact_address": _text(raw.get("emergency_contact_address")) or None,
        }
        end_date = None
        try:
            data["date_of_birth"] = _date(raw.get("date_of_birth"))
        except ValueError as error:
            errors.append(f"Date of Birth: {error}")
        try:
            start_date = _date(raw.get("start_date"))
            if start_date is None:
                errors.append("Start Date: Enter when their current plan started.")
            else:
                data["start_date"] = start_date
        except ValueError as error:
            errors.append(f"Start Date: {error}")
        try:
            end_date = _date(raw.get("end_date"))
            if end_date and data.get("start_date") and end_date < data["start_date"]:
                errors.append("End Date: Must be on or after the start date.")
        except ValueError as error:
            errors.append(f"End Date: {error}")
        try:
            discount = _amount(raw.get("discount"))
            if discount:
                data["discount"] = discount
            paid = _amount(raw.get("amount_paid"))
            if paid:
                method = METHODS.get(_text(raw.get("payment_method")).lower(), "CASH")
                data["payment"] = {"amount": paid, "payment_method": method, "notes": "Imported"}
        except ValueError as error:
            errors.append(str(error))

        if plan is not None:
            data["membership_plan_id"] = plan.id
        if not errors:
            serializer = MemberSerializer(data={k: v for k, v in data.items() if v is not None},
                                          context={"business_account_id": business_account_id})
            if serializer.is_valid():
                data = serializer.validated_data
                paid = data.get("payment", {}).get("amount")
                payable = plan.amount - data.get("discount", Decimal("0"))
                if paid and paid > payable:
                    errors.append(f"Amount Paid: Can't be more than the plan's ₹{payable} (after discount).")
            else:
                errors += _flatten(serializer.errors)

        results.append({
            "row": number, "ok": not errors, "errors": errors, "member_code": code,
            "name": f"{first_name} {last_name}".strip(), "plan": plan.name if plan else plan_name,
            "data": data, "end_date": end_date,
        })
    return results


@transaction.atomic
def import_rows(business_account_id, admin, results):
    """Creates the members from rows that passed check_rows. Returns how many were added."""
    created = 0
    for result in results:
        if not result["ok"]:
            continue
        data = dict(result["data"])
        plan = data.pop("membership_plan_id")
        start_date = data.pop("start_date")
        discount = data.pop("discount", Decimal("0"))
        payment = data.pop("payment", None)
        gym_details = {k: data.pop(k) for k in ("emergency_contact_name", "emergency_contact_phone",
                                                "emergency_contact_address") if k in data}
        member = Member.objects.create(**data, business_account_id=business_account_id,
                                       created_by_business_admin=admin, is_active=True)
        if any(gym_details.values()):
            GymMember.objects.update_or_create(member=member, defaults=gym_details)
        membership = MembershipService.create(member, plan, start_date, discount)
        if result["end_date"]:
            membership.end_date = result["end_date"]
            membership.save(update_fields=["end_date", "updated_at"])
        ActivityService.log(member, "MEMBER_CREATED", f"Imported: {MembershipService.describe(membership)}",
                            admin, membership=membership)
        if payment:
            # Dated on the plan's start day, so old payments don't show up in today's collections.
            payment["payment_date"] = timezone.make_aware(datetime.combine(start_date, time(10, 0)))
            PaymentService.create(member, membership, payment, admin)
        created += 1
    if created:
        transaction.on_commit(lambda: NotificationService.notify_activity(
            business_account_id, admin, "Members imported", f"{admin.name} imported {created} member(s)", None))
    return created
